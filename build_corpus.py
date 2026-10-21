"""把真实业务文档整理成知识库语料（markdown）。

来源三处，全部是真实业务文档，不含客户名单：
  1. 内部项目/WORKFLOW.md        —— 领英抓取工作流（含踩坑表）
  2. 领英海外拓客跟进表.xlsx           —— 「使用说明」「付费判定规则」两个 sheet
  3. 作品集-在不可控平台上建立可靠性.docx —— 按小节拆开

输出到 agent-eval/corpus/，文件名带序号，便于检索命中时定位。
"""
import os
import re

import openpyxl
from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "corpus")

WORKFLOW = r"<path>\WORKFLOW.md"
XLSX = r"<path>\跟进表.xlsx"
DOCX = r"<path>\作品集.docx"


def write(name, title, source, body):
    os.makedirs(OUT, exist_ok=True)
    text = f"# {title}\n\n> 来源：{source}\n\n{body.strip()}\n"
    with open(os.path.join(OUT, name), "w", encoding="utf-8") as f:
        f.write(text)
    print(f"  {name:38s} {len(text.encode('utf-8')) // 1024:>3} KB")


def from_workflow():
    if not os.path.exists(WORKFLOW):
        print("  [!] WORKFLOW.md 不存在，跳过")
        return
    body = open(WORKFLOW, encoding="utf-8").read()
    body = re.sub(r"^# .*\n", "", body, count=1)
    write("01-领英抓取工作流.md", "领英目标职位抓取工作流", "内部项目/WORKFLOW.md", body)


def from_xlsx():
    if not os.path.exists(XLSX):
        print("  [!] 跟进表不存在，跳过")
        return
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    for sheet, name, title in [
        ("使用说明", "02-跟进台账使用说明.md", "跟进台账使用说明"),
        ("付费判定规则", "03-付费判定规则.md", "客户付费判定规则"),
    ]:
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        lines = []
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c not in (None, "")]
            if cells:
                lines.append("| " + " | ".join(cells) + " |" if len(cells) > 1 else cells[0])
        write(name, title, f"领英海外拓客跟进表.xlsx / {sheet}", "\n".join(lines))


SECTION_RE = re.compile(r"^[一二三四五六七八九十]+、")


def from_docx():
    if not os.path.exists(DOCX):
        print("  [!] 案例研究 docx 不存在，跳过")
        return
    doc = Document(DOCX)
    # 标题与首段
    paras = [p.text.strip() for p in doc.paragraphs]
    head_title = next((p for p in paras if p), "在不可控平台上建立可靠性")

    sections, cur = [], None
    for p in paras:
        if not p:
            continue
        if p.startswith("一、") or (SECTION_RE.match(p) and len(p) < 40):
            if cur:
                sections.append(cur)
            cur = {"title": p, "lines": []}
        elif cur is not None:
            cur["lines"].append(p)
    if cur:
        sections.append(cur)

    # 表格按顺序分配给各小节（真实数据表 / 失败分类表 / 演进表 / 迁移表）
    tables = []
    for tb in doc.tables:
        rows = []
        for row in tb.rows:
            rows.append("| " + " | ".join(c.text.strip().replace("\n", " ") for c in row.cells) + " |")
        if rows:
            tables.append("\n".join(rows))

    idx = 4
    for i, sec in enumerate(sections):
        title = sec["title"]
        clean = re.sub(r"^[一二三四五六七八九十]+、", "", title).strip() or f"第 {i + 1} 节"
        body = "\n\n".join(sec["lines"])
        # 把表格塞进最相关的小节
        if "真实数据" in clean and tables:
            body += "\n\n" + tables[0]
        elif "失败分类" in clean and len(tables) > 1:
            body += "\n\n" + tables[1]
        elif "长出来" in clean and len(tables) > 2:
            body += "\n\n" + tables[2]
        elif "迁移" in clean and len(tables) > 3:
            body += "\n\n" + tables[3]
        write(f"{idx:02d}-案例-{clean[:12]}.md", f"{head_title} · {clean}",
              "作品集《在不可控平台上建立可靠性》", body)
        idx += 1


if __name__ == "__main__":
    print("生成业务语料 →", OUT)
    from_workflow()
    from_xlsx()
    from_docx()

    files = sorted(os.listdir(OUT)) if os.path.isdir(OUT) else []
    total = sum(os.path.getsize(os.path.join(OUT, f)) for f in files)
    print(f"\n共 {len(files)} 篇，{total // 1024} KB")
