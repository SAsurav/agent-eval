"""把《领英海外拓客 SOP（精简版）》拆成知识库语料。

拆成 6 篇而不是 1 篇：SOP 太长，整篇塞进去会变成一个巨大的 chunk，
检索命中后返回一大堆无关内容，反而降低回答质量。

脱敏：去掉真实人名、公司名、活动名、酒店名，保留规则与话术结构。
（这些是业务信息，不是公开信息；语料会推到 GitHub。）
"""
import os
import re

from docx import Document

HERE = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.join(HERE, "corpus")

# ===== 原始文件位置：用环境变量指定，仓库里不放真实路径 =====
# 仓库已带生成好的 corpus/，平时不需要跑。
# 需要重跑时：set SOP_DOCX=<...>\SOP.docx
SOP = os.environ.get("SOP_DOCX", r"<path>\SOP.docx")
# ===========================================================

# 脱敏映射：顺序重要，长的先替换
REDACT = [
    ("Daisy Li", "[发送人]"),
    ("Daisy", "[发送人]"),
    ("EnmoreTech", "[公司]"),
    ("Southeast Asia Data Center Liquid Cooling & AI Chip Thermal Summit in Bangkok",
     "[活动全名]"),
    ("Data Center Liquid Cooling & AI Chip Thermal Summit in Bangkok", "[活动全名]"),
    ("Radisson Blu plaze Bangkok", "[活动酒店]"),
    ("this 19-20 Nov", "[活动日期]"),
    ("AWS, Google Cloud, DAMAC Digital, Vertiv, Delta, Schneider Electric and Aurecon",
     "[演讲嘉宾公司列表]"),
]


def clean(text):
    for a, b in REDACT:
        text = text.replace(a, b)
    return text


def write(name, title, body):
    os.makedirs(OUT, exist_ok=True)
    text = f"# {title}\n\n> 来源：领英海外拓客 SOP（精简版）· 已脱敏\n\n{clean(body).strip()}\n"
    with open(os.path.join(OUT, name), "w", encoding="utf-8") as f:
        f.write(text)
    print(f"  {name:38s} {len(text.encode('utf-8')) // 1024:>3} KB")


def table_md(tb):
    rows = []
    for row in tb.rows:
        rows.append("| " + " | ".join(c.text.strip().replace("\n", " ") for c in row.cells) + " |")
    return "\n".join(rows)


def main():
    if not os.path.exists(SOP):
        print(f"  [!] 找不到 {SOP}")
        return
    doc = Document(SOP)
    paras = [p.text.strip() for p in doc.paragraphs]
    tables = [table_md(tb) for tb in doc.tables]

    def find(prefix):
        for i, p in enumerate(paras):
            if p.startswith(prefix):
                return i
        return -1

    def slice_between(start_prefix, end_prefix):
        a = find(start_prefix)
        if a < 0:
            return ""
        b = find(end_prefix) if end_prefix else len(paras)
        if b <= a:
            b = len(paras)
        return "\n".join(p for p in paras[a:b] if p)

    # 1. 日常四步（一、）
    write("12-SOP-日常四步.md", "SOP · 每天照着做的四步",
          slice_between("一、每天照着做", "二、优先级"))

    # 2. 优先级 + 状态（二、三、+ 前两张表）
    body = slice_between("二、优先级", "五、付费判定")
    if len(tables) >= 2:
        body += "\n\n## 优先级判定\n\n" + tables[0] + "\n\n## 状态定义\n\n" + tables[1]
    write("13-SOP-优先级与状态.md", "SOP · 优先级与状态定义", body)

    # 3. 付费判定（SOP 版本，与台账里的版本并存 → 故意保留冲突）
    body = slice_between("五、付费判定", "六、数量红线")
    if len(tables) >= 3:
        body += "\n\n" + tables[2]
    write("17-SOP-付费判定.md", "SOP · 付费判定（按顺序问）", body)

    # 4. 数量红线
    write("14-SOP-数量红线.md", "SOP · 账号操作的数量红线",
          slice_between("六、数量红线", "七、英文话术"))

    # 5. 英文话术
    write("15-SOP-英文话术.md", "SOP · 英文触达话术",
          slice_between("七、英文话术", "八、收工前"))

    # 6. 收工清单
    write("16-SOP-收工清单.md", "SOP · 每日收工检查清单",
          slice_between("八、收工前", None))

    files = sorted(f for f in os.listdir(OUT) if f.endswith(".md"))
    total = sum(os.path.getsize(os.path.join(OUT, f)) for f in files)
    print(f"\n语料共 {len(files)} 篇，{total // 1024} KB")


if __name__ == "__main__":
    main()
